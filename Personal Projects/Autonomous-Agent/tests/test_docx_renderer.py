from pathlib import Path

from docx import Document as DocxDocument

from app.docmodel.models import (
    Block,
    BlockTypeEnum,
    Document,
    Section,
    TableColumn,
    TableData,
)
from app.tools.docx_renderer import render_docx


def _document() -> Document:
    return Document(
        title="Test Document",
        sections=[
            Section(
                title="Overview",
                blocks=[Block(type=BlockTypeEnum.PARAGRAPH, text="Overview text.")],
            ),
            Section(
                title="Budget",
                blocks=[
                    Block(
                        type=BlockTypeEnum.TABLE,
                        table=TableData(
                            columns=[TableColumn(header="Item"), TableColumn(header="Cost")],
                            rows=[["Implementation", "$50,000"]],
                        ),
                    )
                ],
            ),
        ],
    )


def test_render_minimal_document(tmp_path: Path) -> None:
    output_path = render_docx(_document(), tmp_path / "minimal.docx")

    assert output_path.exists()
    assert output_path.stat().st_size > 1000


def test_render_handles_empty_sections(tmp_path: Path) -> None:
    output_path = render_docx(
        Document(title="Empty Sections", sections=[Section(title="Empty")]),
        tmp_path / "empty.docx",
    )

    assert output_path.exists()


def test_render_toc_field_present(tmp_path: Path) -> None:
    output_path = render_docx(_document(), tmp_path / "toc.docx")
    document = DocxDocument(output_path)
    xml = document._part._element.xml

    assert "w:fldSimple" in xml
    assert "TOC" in xml


def test_render_footer_page_numbers(tmp_path: Path) -> None:
    output_path = render_docx(_document(), tmp_path / "footer.docx")
    document = DocxDocument(output_path)
    footer_xml = document.sections[0].footer._element.xml

    assert "PAGE" in footer_xml
    assert "NUMPAGES" in footer_xml
