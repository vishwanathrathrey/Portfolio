"""Deterministic python-docx renderer for internal document models."""

import asyncio
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.config import ensure_output_dir
from app.docmodel.models import BlockTypeEnum, Document, TableData
from app.tools.registry import ToolRegistry, ToolSpec


def render_docx(document: Document, output_path: Path) -> Path:
    """Render an internal Document to a Word document at output_path."""
    docx_doc = DocxDocument()
    _set_margins(docx_doc, inches=1.0)
    _set_default_font(docx_doc, name="Calibri", size_pt=11)

    title = docx_doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(document.title)
    title_run.bold = True
    title_run.font.size = Pt(28)

    if document.subtitle:
        subtitle = docx_doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_run = subtitle.add_run(document.subtitle)
        subtitle_run.italic = True
        subtitle_run.font.size = Pt(14)

    prepared_by = docx_doc.add_paragraph(f"Prepared by: {document.author}")
    prepared_by.alignment = WD_ALIGN_PARAGRAPH.CENTER
    created_at = docx_doc.add_paragraph(document.created_at.strftime("%B %d, %Y"))
    created_at.alignment = WD_ALIGN_PARAGRAPH.CENTER
    docx_doc.add_page_break()

    _add_toc_field(docx_doc)
    docx_doc.add_page_break()

    for section in document.sections:
        docx_doc.add_heading(section.title, level=1)
        for block in section.blocks:
            if block.type is BlockTypeEnum.HEADING1:
                docx_doc.add_heading(block.text or "", level=1)
            elif block.type is BlockTypeEnum.HEADING2:
                docx_doc.add_heading(block.text or "", level=2)
            elif block.type is BlockTypeEnum.PARAGRAPH:
                docx_doc.add_paragraph(block.text or "")
            elif block.type is BlockTypeEnum.BULLET_LIST:
                for item in block.items or []:
                    docx_doc.add_paragraph(item, style="List Bullet")
            elif block.type is BlockTypeEnum.NUMBERED_LIST:
                for item in block.items or []:
                    docx_doc.add_paragraph(item, style="List Number")
            elif block.type is BlockTypeEnum.TABLE and block.table is not None:
                _add_table(docx_doc, block.table)
            elif block.type is BlockTypeEnum.PAGE_BREAK:
                docx_doc.add_page_break()

    _add_footer_page_numbers(docx_doc)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    docx_doc.save(output_path)
    return output_path


def _set_margins(docx_doc: Any, inches: float) -> None:
    for section in docx_doc.sections:
        section.top_margin = Inches(inches)
        section.bottom_margin = Inches(inches)
        section.left_margin = Inches(inches)
        section.right_margin = Inches(inches)


def _set_default_font(docx_doc: Any, name: str, size_pt: int) -> None:
    normal_style = docx_doc.styles["Normal"]
    normal_style.font.name = name
    normal_style.font.size = Pt(size_pt)
    normal_style.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _add_toc_field(docx_doc: Any) -> None:
    docx_doc.add_heading("Table of Contents", level=1)
    paragraph = docx_doc.add_paragraph()
    toc_field = OxmlElement("w:fldSimple")
    toc_field.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
    paragraph._p.append(toc_field)
    docx_doc.add_paragraph("(Right-click -> Update Field to populate the table of contents.)")


def _add_footer_page_numbers(docx_doc: Any) -> None:
    for section in docx_doc.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run("Page ")
        _append_field(paragraph, "PAGE")
        paragraph.add_run(" of ")
        _append_field(paragraph, "NUMPAGES")


def _append_field(paragraph: Any, instruction: str) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    paragraph._p.append(field)


def _shade_cell(cell: Any, fill_hex: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell_properties.append(shading)


def _add_table(docx_doc: Any, table_data: TableData) -> None:
    table = docx_doc.add_table(rows=1, cols=len(table_data.columns))
    table.autofit = True
    header_cells = table.rows[0].cells
    for index, column in enumerate(table_data.columns):
        header_cell = header_cells[index]
        header_cell.text = column.header
        header_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _shade_cell(header_cell, "D9E2F3")
        for run in header_cell.paragraphs[0].runs:
            run.bold = True

    for row in table_data.rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _apply_column_alignments(table, table_data)


def _apply_column_alignments(table: Any, table_data: TableData) -> None:
    alignments = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    for row in table.rows:
        for index, column in enumerate(table_data.columns):
            for paragraph in row.cells[index].paragraphs:
                paragraph.alignment = alignments[column.align]


async def render_docx_tool(document: Document, output_path: str) -> str:
    """Render a document in a worker thread for async registry dispatch."""
    rendered_path = await asyncio.to_thread(
        render_docx,
        document,
        _resolve_output_path(output_path),
    )
    return str(rendered_path)


def _resolve_output_path(output_path: str) -> Path:
    path = Path(output_path)
    if path.parent == Path("."):
        return ensure_output_dir() / path
    return path


def register_tool(registry: ToolRegistry) -> None:
    """Register the DOCX-rendering tool."""
    registry.register(
        ToolSpec(
            name="render_docx",
            description="Render the assembled Document to a .docx file.",
            args_schema={"document": "Document", "output_path": "str"},
            handler=render_docx_tool,
        )
    )
